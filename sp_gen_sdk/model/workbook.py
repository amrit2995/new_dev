import openpyxl
from pydantic import BaseModel
from enum import Enum

class Spreadsheet:
    def __init__(self, worksheet):
        self.worksheet = worksheet

    def get_cell(self, row, column):
        return self.worksheet.cell(row=row, column=column).value

    def get_all_cols(self):
        cols = []
        for col in self.worksheet.iter_cols(values_only=True):
            cols.append(col[0])  # Assuming the first row contains column names
        return cols

    def iter_rows(self):
        first_row = True
        for row in self.worksheet.iter_rows(values_only=True):
            if first_row:
                first_row = False
                continue
            yield row

    @property
    def title(self):
        return self.worksheet.title
    
    def get_blue_rows(self):
        blue_rows = []
        for row in self.worksheet.iter_rows():
            # Check if any cell in the row has a blue fill
            for cell in row:
                fill = cell.fill
                if fill and fill.start_color and fill.start_color.type == 'rgb':
                    # Common blue RGB hex codes (e.g., 'FF0000FF' for blue)
                    if fill.start_color.rgb in ['FF0000FF', 'FF1F497D', 'FF0070C0']:
                        blue_rows.append([c.value for c in row])
                        break
        return blue_rows

class Workbook:
    def __init__(self, file_path):
        self.file_path = file_path
        self.workbook = openpyxl.load_workbook(file_path)
        self.sheets = {ws.title: Spreadsheet(ws) for ws in self.workbook.worksheets}

    # @property
    # def name(self):
    #     import pdb;pdb.set_trace()
    #     return self.workbook.title

    # @property
    # def sheets(self, name):
    #     return self.sheets.get(name)

    @property
    def sheet_names_list(self):
        return self.workbook.sheetnames
    
    def get_sheet_by_name(self, name):
        return self.sheets[name]

    def close(self):
        self.workbook.close()